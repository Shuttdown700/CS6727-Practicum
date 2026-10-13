#!/usr/bin/env python3

from __future__ import annotations

import re
import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

import openpyxl
from openpyxl.utils.cell import range_boundaries

### PATHS
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
WORKBOOK = SCRIPT_DIR / "input" / "Control Traceability Matrix.xlsx"
OUTPUT_DIR = SCRIPT_DIR / "output"

# WORKBOOK SHEET NAMES
SHEET_TREE = "Attack Tree Structure"
SHEET_LEAF = "Attack Leaf Risk Matrix"
SHEET_EDGES = "Attack Tree Edges"

# GLOBAL INFOGRAPHIC SETTINGS
DPI = 200
PLOT_STYLE = "ggplot"

# SOLUTION VARIANTS & ASSOCIATED COLUMN NAMES
VARIANTS = [
    ("Baseline",      "Baseline Status",      "Baseline Likelihood",      "Baseline Residual Risk Score"),
    ("Observational", "Observational Status",  "Observational Likelihood", "Observational Residual Risk Score"),
    ("Mediated",      "Mediated Status",       "Mediated Likelihood",      "Mediated Residual Risk"),
]
BASELINE = VARIANTS[0][0]

# COLUMN NAMES: EDGE SHEET
E_TREE        = "Tree"
E_PARENT      = "Parent ID"
E_CHILD       = "Child ID"
E_ORDER       = "Child Order"
EDGE_COLUMNS = [E_TREE, 
                E_PARENT, 
                E_CHILD,
                E_ORDER]


# COLUMN NAMES: LEAF SHEET
C_ID = "ID"
C_DESC = "Attack Leaf Description"
C_TYPE = "Attack Leaf Type"
C_TREES = "Attack Trees"
C_RISKCLASS = "Risk classes"
C_CONDTYPE = "Conditional Type"
C_VAL_FREQ = "Value | Frequency"
C_VAL_LOAI = "Value | Likelihood of Adverse Impact"
C_VAL_OL = "Value | Overall Likelihood"
C_VAL_IMPACT = "Value | Impact"
C_FEATURE = "Solution Feature"
C_CONFIDENCE = "[T] Solution Effect Confidence"

# MEDIATION LIKELIHOOD SEMANTIC MAPPINGS
MEDIATION_STATUS_LIKELIHOOD = {
    "Unaddressed": 10,
    "Partially Addressed": 5,
    "Fully Addressed": 0
}
MEDIATION_STATUS_ORDER = ["Unaddressed", 
                "Partially Addressed", 
                "Fully Addressed"]
MEDIATION_STATUS_CONSTANT = "Constant"
MEDIATION_STATUS_NA = "N/A"


# NIST SP 800-30 (Rev. 1) SEMI-QUANTITATIVE SCALES
BANDS = [
    ("Very Low",   0),
    ("Low",        2),
    ("Moderate",   5),
    ("High",       8),
    ("Very High", 10),
]
BAND_NAME = {v: n for n, v in BANDS}
BAND_VALUE = {n: v for n, v in BANDS}
BAND_LADDER = [v for _, v in BANDS]

# Table G-5: Overall Likelihood: [frequency][likelihood_of_adverse_impact]
TABLE_G5_OVERALL_LIKELIHOOD = {
    10: {0: 2, 2: 5, 5: 8, 8: 10, 10: 10},
    8:  {0: 2, 2: 5, 5: 5, 8: 8,  10: 10},
    5:  {0: 2, 2: 2, 5: 5, 8: 5,  10: 8},
    2:  {0: 0, 2: 2, 5: 2, 8: 5,  10: 5},
    0:  {0: 0, 2: 0, 5: 2, 8: 2,  10: 5},
}

# Table I-2: Level of Risk: [likelihood][impact]
TABLE_I2_LEVEL_OF_RISK = {
    10: {0: 0, 2: 2, 5: 5, 8: 8, 10: 10},
    8:  {0: 0, 2: 2, 5: 5, 8: 8, 10: 10},
    5:  {0: 0, 2: 2, 5: 5, 8: 5, 10: 8},
    2:  {0: 0, 2: 2, 5: 2, 8: 2, 10: 5},
    0:  {0: 0, 2: 0, 5: 0, 8: 2, 10: 2},
}

# COLORS
SERIES = {
    "Baseline": "#2a78d6",
    "Observational": "#eb6834",
    "Mediated":"#1baf7a",
}

DIV_GOOD = "#2a78d6"
DIV_BAD  = "#e34948"
DIV_ZERO = "#9b9a94"

SEQ_BLUE = ["#86b6ef",
            "#5598e7",
            "#2a78d6",
            "#1c5cab",
            "#104281"]

INK = "#0b0b0b"
INK_SOFT = "#52514e"
INK_MUTED = "#6e6d68"

MEDIATION_STATUS_FILL = {
    "Unaddressed":         "#104281",
    "Partially Addressed": "#5598e7",
    "Fully Addressed":     "#cde2fb",
}

# UTILITY FUNCTIONS
def normalize_canonical(header) -> str:
    """Normalize a header for canonical comparison."""
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def assess_overall_likelihood(freq: int, loai: int) -> int:
    """Return the overall likelihood of adverse impact given frequency and likelihood of adverse impact."""
    return TABLE_G5_OVERALL_LIKELIHOOD[freq][loai]


def assess_risk(likelihood: int, impact: int) -> int:
    """Return the level of risk given likelihood and impact."""
    return TABLE_I2_LEVEL_OF_RISK[likelihood][impact]


def shift_band(value: int, steps: int) -> int:
    """Shift a band value up or down the ladder by `steps`."""
    i = BAND_LADDER.index(value)
    return BAND_LADDER[max(0, min(len(BAND_LADDER) - 1, i + steps))]


# APPLY MATPLOTLIB STYLE AND PARAMS
def apply_style() -> None:
    """Apply the global matplotlib style and parameters."""
    plt.style.use(PLOT_STYLE)
    plt.rcParams.update({
        "figure.dpi": DPI,
        "savefig.dpi": DPI,
        "savefig.bbox": "tight",
        "font.size": 9,
        "axes.titlesize": 12,
        "axes.titleweight": "bold",
        "axes.titlepad": 10,
        "axes.labelsize": 9.5,
        "axes.labelcolor": INK_SOFT,
        "axes.edgecolor": "#cfcec8",
        "axes.titlecolor": INK,
        "xtick.color": INK_SOFT,
        "ytick.color": INK_SOFT,
        "xtick.labelsize": 8.5,
        "ytick.labelsize": 8.5,
        "legend.fontsize": 8.5,
        "legend.frameon": False,
        "grid.color": "#ffffff",
        "grid.linewidth": 0.8,
        "lines.linewidth": 2.0,
        "lines.markersize": 8,
        "axes.prop_cycle": matplotlib.cycler(
            color=[SERIES["Baseline"], SERIES["Observational"], SERIES["Mediated"]]
        ),
    })


def finish(fig, ax_or_axes, path: Path, note: str | None = None) -> Path:
    """Attach an optional footnote, save, close. Returns the written path."""
    if note:
        fig.text(0.005, -(0.55 / fig.get_figheight()), note, ha="left", va="top",
                 fontsize=7.2, color=INK_MUTED, wrap=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(path)
    plt.close(fig)
    return path


def _load_sheet_rows(wb, sheet_name: str):
    """Load a worksheet and return (header, body) as lists of tuples."""
    if sheet_name not in wb.sheetnames:
        return None, None
    ws = wb[sheet_name]
    ref = None
    try:
        tables = list(ws.tables.items())
        if len(tables) == 1:
            ref = tables[0][1]
        elif tables:
            key = re.sub(r"\W+", "", sheet_name).lower()
            match = [r for n, r in tables if re.sub(r"\W+", "", n).lower() == key]
            if match:
                ref = match[0]
            else:
                def area(r):
                    c0, r0, c1, r1 = range_boundaries(r)
                    return (c1 - c0 + 1) * (r1 - r0 + 1)
                ref = max((r for _, r in tables), key=area)
    except Exception:
        ref = None
    cells = list(ws[ref]) if ref else list(ws.iter_rows())
    rows = [tuple(c.value for c in row) for row in cells]
    if not rows:
        return None, None
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    body = [r for r in rows[1:] if any(v is not None and str(v).strip() != "" for v in r)]
    return header, body

# HELPER FUNCTIONS
def _normalize_cell(v):
    """Normalise a cell: trailing spaces stripped, blanks and 'N/A' to None."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v == "" or v.upper() == "N/A":
            return None
    return v


def _split_list(v) -> list[str]:
    """Split a comma-separated cell into a list of stripped strings."""
    if v is None:
        return []
    return [p.strip() for p in str(v).split(",") if p.strip()]

# DATA CLASSES
class Leaf:
    """One row of the Attack Leaf Risk Matrix."""

    __slots__ = ("id", "desc", "type", "trees", "risk_classes", "cond_type",
                 "freq", "loai", "ol", "impact", "feature", "confidence",
                 "status", "likelihood", "risk")

    def __init__(self, rec: dict):
        self.id = str(rec[C_ID]).strip()
        self.desc = str(rec.get(C_DESC) or "").strip()
        self.type = str(rec.get(C_TYPE) or "").strip()
        self.trees = _split_list(rec.get(C_TREES))
        self.risk_classes = _split_list(rec.get(C_RISKCLASS))
        self.cond_type = _normalize_cell(rec.get(C_CONDTYPE))
        self.freq = _normalize_cell(rec.get(C_VAL_FREQ))
        self.loai = _normalize_cell(rec.get(C_VAL_LOAI))
        self.ol = _normalize_cell(rec.get(C_VAL_OL))
        self.impact = _normalize_cell(rec.get(C_VAL_IMPACT))
        feat = _normalize_cell(rec.get(C_FEATURE))
        self.feature = None if feat in (None, "None") else str(feat).strip()
        self.confidence = _normalize_cell(rec.get(C_CONFIDENCE))
        self.status, self.likelihood, self.risk = {}, {}, {}
        for name, scol, lcol, rcol in VARIANTS: # status, likelihood, risk
            st = rec.get(scol)
            st = str(st).strip() if st is not None else None
            if st in (None, "", MEDIATION_STATUS_NA):
                self.status[name] = None
                self.likelihood[name] = None
                self.risk[name] = None
                continue
            self.status[name] = st
            self.likelihood[name] = _normalize_cell(rec.get(lcol))
            self.risk[name] = _normalize_cell(rec.get(rcol))

    @property
    def is_condition(self) -> bool:
        """Return True if this leaf is a conditional risk, False otherwise."""
        return self.type.lower().startswith("cond")

    @property
    def is_frequency(self) -> bool:
        """Return True if this leaf is a frequency risk, False otherwise."""
        return self.type.lower().startswith("freq")

    def present(self, variant: str) -> bool:
        """Return True if this leaf has a risk value for the given variant, False otherwise."""
        return self.risk.get(variant) is not None

    def short(self, width: int = 46) -> str:
        """Return a shortened version of the description, suitable for plotting."""
        return textwrap.shorten(self.desc, width=width, placeholder="...")

    def __repr__(self):
        return f"<Leaf {self.id} {self.type}>"


class Node:
    """One row of the Attack Tree Structure sheet."""

    __slots__ = ("id", "node_type", "gate", "label")

    def __init__(self, rec: dict):
        self.id = str(rec["Node ID"]).strip()
        self.node_type = str(rec.get("Node Type") or "").strip()
        self.gate = _normalize_cell(rec.get("Gate Type"))
        self.label = str(rec.get("Label") or "").strip()

    @property
    def is_root(self) -> bool:
        """Return True if this node is a root node, False otherwise."""
        return self.node_type.lower() == "root"

    @property
    def is_gate(self) -> bool:
        """Return True if this node is a gate node, False otherwise."""
        return self.node_type.lower() == "gate"

    @property
    def is_leaf(self) -> bool:
        """Return True if this node is a leaf node, False otherwise."""
        return self.node_type.lower() == "leaf"

    def __repr__(self):
        return f"<Node {self.id} {self.node_type} {self.gate}>"


class Model:
    """The loaded workbook: nodes, leaves, and optional edges."""

    def __init__(self, workbook: Path = WORKBOOK):
        wb = openpyxl.load_workbook(workbook, data_only=True)
        self.workbook_path = workbook

        hdr, body = _load_sheet_rows(wb, SHEET_LEAF)
        if hdr is None:
            raise SystemExit(f"worksheet '{SHEET_LEAF}' not found in {workbook}")
        self.leaves = [Leaf(dict(zip(hdr, r))) for r in body if r and r[0]] # skip blank rows
        self.by_id = {lf.id: lf for lf in self.leaves}

        hdr, body = _load_sheet_rows(wb, SHEET_TREE)
        if hdr is None:
            raise SystemExit(f"worksheet '{SHEET_TREE}' not found in {workbook}")
        self.nodes = [Node(dict(zip(hdr, r))) for r in body if r and r[0]] # skip blank rows
        self.node_by_id = {n.id: n for n in self.nodes}
        self.roots = [n.id for n in self.nodes if n.is_root]

        self.trees = list(self.roots) # all trees are roots, but some leaves may reference additional trees
        for lf in self.leaves:
            for t in lf.trees:
                if t not in self.trees:
                    self.trees.append(t)

        self.edges = []
        hdr, body = _load_sheet_rows(wb, SHEET_EDGES)
        if hdr is not None:
            cols = {normalize_canonical(h): i for i, h in enumerate(hdr) if normalize_canonical(h)}
            need = {normalize_canonical(E_TREE): E_TREE,
                    normalize_canonical(E_PARENT): E_PARENT,
                    normalize_canonical(E_CHILD): E_CHILD}
            missing = [label for key, label in need.items() if key not in cols]
            if missing:
                raise SystemExit(
                    f"\nworksheet '{SHEET_EDGES}' is present but missing required "
                    f"column(s): {', '.join(missing)}\n\n"
                    f"  workbook: {workbook}\n"
                    f"  found headers: {', '.join(str(h) for h in hdr if h)}\n\n"
                    f"Expected: {' | '.join(EDGE_COLUMNS)}   ({E_ORDER} is optional)\n"
                    "Matching ignores case, spaces, underscores and hyphens, so\n"
                    "'Parent ID', 'parent_id' and 'PARENT-ID' are all accepted.\n"
                )
            ci_tree = cols[normalize_canonical(E_TREE)]
            ci_parent = cols[normalize_canonical(E_PARENT)]
            ci_child = cols[normalize_canonical(E_CHILD)]
            ci_order = cols.get(normalize_canonical(E_ORDER))
            for i, r in enumerate(body):
                tree = _normalize_cell(r[ci_tree])
                parent = _normalize_cell(r[ci_parent])
                child = _normalize_cell(r[ci_child])
                if not (tree and parent and child):
                    continue
                order = r[ci_order] if ci_order is not None else None
                self.edges.append({
                    "tree": str(tree).strip(),
                    "parent_id": str(parent).strip(),
                    "child_id": str(child).strip(),
                    "child_order": order if isinstance(order, (int, float)) else i,
                })

        self.children = {}
        for e in self.edges:
            self.children.setdefault((e["tree"], e["parent_id"]), []).append(e)
        for key in self.children:
            self.children[key].sort(key=lambda e: e["child_order"])

        self._check_cached_values()

    def _check_cached_values(self) -> None:
        """Check that the workbook has cached values for the scoring columns."""
        if not self.leaves:
            raise SystemExit(f"no leaf rows parsed from '{SHEET_LEAF}' in {self.workbook_path}")
        blank = [lf.id for lf in self.leaves if lf.impact is None]
        if len(blank) <= len(self.leaves) // 3:
            return
        raise SystemExit(
            f"\n{len(blank)} of {len(self.leaves)} leaves have no impact value, and the per-variant\n"
            f"risk columns are almost certainly empty too.\n\n"
            f"  workbook: {self.workbook_path}\n\n"
            "The scoring columns are Excel formulas, and this loader reads the cached\n"
            "result stored beside each one. That cache is dropped by any tool that saves\n"
            "the file without recalculating -- openpyxl does exactly this. The formulas\n"
            "survive; their values do not.\n\n"
            "To fix: open the workbook in Excel (or LibreOffice), let it recalculate, and\n"
            "save. Never write to this workbook with openpyxl.\n\n"
            "If you are adding the 'Attack Tree Edges' sheet, add it in Excel for the same\n"
            "reason -- a script that appends the sheet will blank the whole model.\n"
        )


    @property
    def has_edges(self) -> bool:
        """Return True if the optional edge table is present, False otherwise."""
        return bool(self.edges)


    @property
    def variant_names(self) -> list[str]:
        """Return the list of variant names in order."""
        return [v[0] for v in VARIANTS]


    def leaves_in(self, tree: str) -> list[Leaf]:
        """Return the list of leaves that reference `tree` in their `Attack Trees` field."""
        return [lf for lf in self.leaves if tree in lf.trees]


    def common_leaves(self) -> list[Leaf]:
        """Leaves scored under every variant -- the only like-for-like set."""
        return [lf for lf in self.leaves
                if all(lf.present(v) for v in self.variant_names)]


    def variant_only_leaves(self, variant: str) -> list[Leaf]:
        """Leaves that exist under `variant` but not under the baseline."""
        return [lf for lf in self.leaves
                if lf.present(variant) and not lf.present(BASELINE)]


    def total(self, variant: str, leaves=None) -> int:
        """Return the total risk for a variant."""
        pool = self.leaves if leaves is None else leaves
        return sum(lf.risk[variant] for lf in pool if lf.present(variant))


    def tree_total(self, tree: str, variant: str) -> int:
        """Return the total risk for a variant under a tree."""
        return sum(lf.risk[variant] for lf in self.leaves_in(tree)
                   if lf.present(variant))


    def propagate(self, tree: str, variant: str) -> dict[str, int | None]:
        """Propagate the leaf risk values up the tree, returning a dict of node_id -> value."""
        if not self.has_edges:
            raise RuntimeError("no edge table loaded; cannot propagate")
        memo: dict[str, int | None] = {}

        def walk(node_id: str, seen: frozenset) -> int | None:
            """Recursively walk the tree, returning the propagated value for `node_id`."""
            if node_id in memo:
                return memo[node_id]
            if node_id in seen:
                raise SystemExit(f"cycle in tree {tree} at {node_id}")
            kids = self.children.get((tree, node_id), [])
            if not kids:
                lf = self.by_id.get(node_id)
                val = lf.risk[variant] if lf and lf.present(variant) else None
                memo[node_id] = val
                return val
            vals = [v for v in (walk(k["child_id"], seen | {node_id}) for k in kids)
                    if v is not None]
            if not vals:
                memo[node_id] = None
                return None
            gate = (self.node_by_id[node_id].gate or "OR").upper() \
                if node_id in self.node_by_id else "OR"
            memo[node_id] = max(vals) if gate == "OR" else min(vals)
            return memo[node_id]

        walk(tree, frozenset())
        return memo

    def critical_child(self, tree: str, variant: str) -> tuple[str, int] | None:
        """Which root child sets the root value -- the free prioritisation signal."""
        if not self.has_edges:
            return None
        vals = self.propagate(tree, variant)
        kids = self.children.get((tree, tree), [])
        scored = [(k["child_id"], vals.get(k["child_id"])) for k in kids]
        scored = [(c, v) for c, v in scored if v is not None]
        if not scored:
            return None
        gate = (self.node_by_id[tree].gate or "OR").upper()
        return max(scored, key=lambda x: x[1]) if gate == "OR" else min(scored, key=lambda x: x[1])
